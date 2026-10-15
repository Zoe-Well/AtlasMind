"""Bounded parser port and adapters for supported enterprise file types."""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass
from io import BytesIO
from typing import Any, Protocol

from bs4 import BeautifulSoup, Comment, NavigableString, Tag
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from markdown_it import MarkdownIt
from pypdf import PdfReader


class ParsingFailed(ValueError):
    def __init__(self, reason: str) -> None:
        self.reason = reason
        super().__init__(reason)


@dataclass(frozen=True, slots=True)
class DocumentBlock:
    text: str
    kind: str = "p"
    heading_path: tuple[str, ...] = ()
    source: str = ""


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    paragraphs: tuple[str, ...]
    language: str
    blocks: tuple[DocumentBlock, ...] = ()


class Parser(Protocol):
    def parse(self, content: bytes) -> ParsedDocument: ...


def _normalize(parts: list[str]) -> ParsedDocument:
    paragraphs = tuple(" ".join(part.split()) for part in parts if part and part.strip())
    if not paragraphs:
        raise ParsingFailed("NO_EXTRACTABLE_TEXT")
    joined = "".join(paragraphs)
    cjk = sum("\u4e00" <= char <= "\u9fff" for char in joined)
    return ParsedDocument(paragraphs=paragraphs, language="zh" if cjk else "en")


def _document(blocks: list[DocumentBlock]) -> ParsedDocument:
    clean = tuple(
        DocumentBlock(
            " ".join(b.text.split()),
            b.kind,
            tuple(" ".join(h.split()) for h in b.heading_path),
            " ".join(b.source.split()),
        )
        for b in blocks
        if b.text.strip()
    )
    normalized = _normalize([b.text for b in clean])
    return ParsedDocument(normalized.paragraphs, normalized.language, clean)


def _decode(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ParsingFailed("TEXT_ENCODING_UNSUPPORTED") from exc


def _join_lines(lines: list[str]) -> str:
    """Join wrapped prose without introducing spaces inside Chinese words."""
    text = ""
    for line in lines:
        line = line.strip()
        separator = " " if text and not re.search(r"[\u3400-\u9fff]$", text) else ""
        text += separator + line
    return text


def _prose(text: str) -> list[str]:
    """Blank lines delimit prose; lists and Q/A starts remain semantic boundaries."""
    result: list[str] = []
    current: list[str] = []
    for line in text.splitlines():
        boundary = re.match(r"^\s*(?:[-*•]|\d+[.)、]|[问答QA][:\uff1a])\s*", line)
        # Keep an answer with its immediately preceding question.
        answer = re.match(r"^\s*[答A][:\uff1a]", line)
        question = current and re.match(r"^\s*[问Q][:\uff1a]", current[0])
        if current and (not line.strip() or (boundary and not (answer and question))):
            result.append(_join_lines(current))
            current = []
        if line.strip():
            current.append(line)
    if current:
        result.append(_join_lines(current))
    # A blank line between Q and A should not separate the rule.
    paired: list[str] = []
    for part in result:
        if (
            re.match(r"^[答A][:\uff1a]", part)
            and paired
            and re.match(r"^[问Q][:\uff1a]", paired[-1])
        ):
            paired[-1] += " " + part
        else:
            paired.append(part)
    return paired


def _table_rows(
    rows: list[list[str]], path: tuple[str, ...], source: str, *, has_header: bool = True
) -> list[DocumentBlock]:
    if not rows or not any(value.strip() for row in rows for value in row):
        return []
    # Preserve the first row too: DOCX does not always mark a header explicitly.
    headers = rows[0]
    label = "表头\uff1a" if has_header else "表首行\uff1a"
    path = (*path, label + " | ".join(headers)) if any(headers) else path
    result = [DocumentBlock(" | ".join(headers), "table", path, source + "\uff0c行 1")]
    for number, row in enumerate(rows[1:], 2):
        cells = []
        for i, value in enumerate(row):
            column = headers[i] if has_header and i < len(headers) and headers[i] else f"列 {i + 1}"
            cells.append(f"{column}\uff1a{value}")
        result.append(
            DocumentBlock(" | ".join(cells), "table", path, source + f"\uff0c行 {number}")
        )
    return result


class TextParser:
    def parse(self, content: bytes) -> ParsedDocument:
        blocks: list[DocumentBlock] = []
        headings: list[tuple[int, str]] = []
        lines: list[str] = []

        def flush() -> None:
            blocks.extend(
                DocumentBlock(part, "p", tuple(h[1] for h in headings))
                for part in _prose("\n".join(lines))
            )
            lines.clear()

        for line in _decode(content).splitlines():
            match = re.match(r"^(#{1,6})\s+(.+)$", line.strip())
            chapter = re.match(r"^第[一二三四五六七八九十百\d]+章\s*.+$", line.strip())
            if match or chapter:
                flush()
                level, title = (len(match[1]), match[2]) if match else (1, line.strip())
                while headings and headings[-1][0] >= level:
                    headings.pop()
                headings.append((level, title))
                blocks.append(DocumentBlock(title, f"h{level}", tuple(h[1] for h in headings)))
            else:
                lines.append(line)
        flush()
        return _document(blocks)


class MarkdownParser:
    def parse(self, content: bytes) -> ParsedDocument:
        soup = BeautifulSoup(MarkdownIt().enable("table").render(_decode(content)), "html.parser")
        for element in soup(["script", "style", "noscript", "template"]):
            element.decompose()
        blocks: list[DocumentBlock] = []
        headings: list[tuple[int, str]] = []
        table_number = 0

        def visit(node: Tag) -> None:
            nonlocal table_number
            name = node.name
            value = " ".join(node.get_text().split())
            if re.fullmatch(r"h[1-6]", name):
                level = int(name[1])
                while headings and headings[-1][0] >= level:
                    headings.pop()
                headings.append((level, value))
                blocks.append(DocumentBlock(value, name, tuple(h[1] for h in headings)))
            elif name == "table":
                table_number += 1
                rows = [
                    [
                        " ".join(c.get_text().split())
                        for c in row.find_all(["th", "td"], recursive=False)
                    ]
                    for row in node.find_all("tr")
                ]
                blocks.extend(
                    _table_rows(rows, tuple(h[1] for h in headings), f"表 {table_number}")
                )
            elif name in {"p", "pre", "li"}:
                # Nested lists are visited separately, without duplicating their text.
                if name == "li":
                    value = " ".join(
                        "".join(
                            str(c) if not isinstance(c, Tag) else c.get_text()
                            for c in node.children
                            if not isinstance(c, Comment)
                            and (not isinstance(c, Tag) or c.name not in {"ul", "ol"})
                        ).split()
                    )
                kind = name
                if name == "li" and isinstance(node.parent, Tag) and node.parent.name == "ol":
                    start = str(node.parent.get("start", "1"))
                    first = int(start) if re.fullmatch(r"\d{1,6}", start) else 1
                    number = first + len(node.find_previous_siblings("li"))
                    value, kind = f"{number}. {value}", "p"
                blocks.append(DocumentBlock(value, kind, tuple(h[1] for h in headings)))
                if name == "li":
                    for child in node.find_all(["ul", "ol"], recursive=False):
                        visit(child)
            else:
                if not node.find(["p", "pre", "li", "table", "h1", "h2", "h3", "h4", "h5", "h6"]):
                    blocks.append(DocumentBlock(value, "p", tuple(h[1] for h in headings)))
                    return
                for child in node.children:
                    if isinstance(child, Tag):
                        visit(child)
                    elif isinstance(child, NavigableString) and not isinstance(child, Comment):
                        blocks.append(DocumentBlock(str(child), "p", tuple(h[1] for h in headings)))

        for node in soup.children:
            if isinstance(node, Tag):
                visit(node)
        # Keep adjacent explicit Markdown Q/A paragraphs in the same atomic block.
        paired: list[DocumentBlock] = []
        for block in blocks:
            if (
                paired
                and block.kind == paired[-1].kind == "p"
                and block.heading_path == paired[-1].heading_path
                and re.match(r"^[答A][:\uff1a]", block.text)
                and re.match(r"^[问Q][:\uff1a]", paired[-1].text)
            ):
                before = paired.pop()
                paired.append(
                    DocumentBlock(before.text + " " + block.text, "p", block.heading_path)
                )
            else:
                paired.append(block)
        return _document(paired)


class HtmlParser:
    def parse(self, content: bytes) -> ParsedDocument:
        soup = BeautifulSoup(content, "html.parser")
        for element in soup(["script", "style", "noscript"]):
            element.decompose()
        return _normalize(soup.get_text("\n").splitlines())


def _pdf_blocks(cleaned: list[str], number: int) -> list[DocumentBlock]:
    blocks: list[DocumentBlock] = []
    # Keep aligned rows together with their first row. No inferred PDF headers.
    prose_lines: list[str] = []
    column_rows: list[list[str]] = []
    block_number = 0

    def flush_prose() -> None:
        nonlocal block_number
        for part in _prose("\n".join(prose_lines)):
            block_number += 1
            blocks.append(
                DocumentBlock(part, "p", (), f"PDF 第 {number} 页\uff0c文本块 {block_number}")
            )
        prose_lines.clear()

    def flush_columns() -> None:
        nonlocal block_number
        if column_rows:
            block_number += 1
            blocks.extend(
                _table_rows(
                    column_rows,
                    (),
                    f"PDF 第 {number} 页\uff0c分栏块 {block_number}",
                    has_header=False,
                )
            )
            column_rows.clear()

    empty_lines = 0
    for row in cleaned:
        if " | " in row:
            flush_prose()
            column_rows.append(row.split(" | "))
            empty_lines = 0
        elif not row.strip() and column_rows:
            empty_lines += 1
            if empty_lines >= 2:
                flush_columns()
        else:
            flush_columns()
            prose_lines.append(row)
    flush_columns()
    flush_prose()
    return blocks


class PdfParser:
    def parse(self, content: bytes) -> ParsedDocument:
        try:
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise ParsingFailed("ENCRYPTED_FILE")
            pages: list[list[str]] = []
            margins: list[set[str]] = []
            for page in reader.pages:
                edge: set[str] = set()
                bottom, top = float(page.mediabox.bottom), float(page.mediabox.top)
                height = top - bottom

                def visitor(
                    text: str,
                    cm: list[float],
                    tm: list[float],
                    font: dict[str, Any] | None,
                    size: float,
                    *,
                    edge: set[str] = edge,
                    bottom: float = bottom,
                    top: float = top,
                    height: float = height,
                ) -> None:
                    y = tm[4] * cm[1] + tm[5] * cm[3] + cm[5]
                    if y < bottom + height * 0.07 or y > top - height * 0.07:
                        edge.update(
                            " ".join(line.split()) for line in text.splitlines() if line.strip()
                        )

                plain = page.extract_text(visitor_text=visitor) or ""
                if not plain.strip() and page.images:
                    raise ParsingFailed("PDF_OCR_REQUIRED")
                layout = (
                    page.extract_text(extraction_mode="layout", layout_mode_strip_rotated=False)
                    or plain
                )
                pages.append(layout.splitlines())
                margins.append(edge)
            counts = Counter(line for edge in margins for line in edge)
            repeated = {
                line for line, count in counts.items() if count >= max(3, (len(pages) + 1) // 2)
            }
            blocks: list[DocumentBlock] = []
            for number, lines in enumerate(pages, 1):
                cleaned = []
                occupied = [i for i, line in enumerate(lines) if line.strip()]
                outer_lines = set(occupied[:1] + occupied[-1:])
                for line_index, line in enumerate(lines):
                    normalized = " ".join(line.split())
                    page_number = re.fullmatch(r"(?:第\s*)?(\d+)\s*(?:页|/\s*\d+)?", normalized)
                    if (
                        line_index in outer_lines
                        and normalized in margins[number - 1]
                        and (
                            normalized in repeated
                            or (page_number and int(page_number[1]) == number)
                        )
                    ):
                        continue
                    # Preserve visible column separators instead of flattening table rows.
                    row = re.sub(r"\S[ \t]{3,}(?=\S)", lambda m: m[0][0] + " | ", line.strip())
                    cleaned.append(row)
                blocks.extend(_pdf_blocks(cleaned, number))
            return _document(blocks)
        except ParsingFailed:
            raise
        except Exception as exc:
            raise ParsingFailed("PDF_PARSE_FAILED") from exc


class DocxParser:
    def parse(self, content: bytes) -> ParsedDocument:
        try:
            document = Document(BytesIO(content))
            blocks: list[DocumentBlock] = []
            headings: list[tuple[int, str]] = []
            table_number = 0

            def table_rows(table: Table) -> list[list[str]]:
                rows = []
                for row in table.rows:
                    cells = []
                    for cell in row.cells:
                        parts = []
                        for item in cell.iter_inner_content():
                            if isinstance(item, Paragraph):
                                parts.append(item.text)
                            else:
                                parts.extend(" | ".join(r) for r in table_rows(item))
                        cells.append(" ".join(" ".join(parts).split()))
                    rows.append(cells)
                return rows

            for item in document.iter_inner_content():
                if isinstance(item, Table):
                    table_number += 1
                    header_flags = (
                        item.rows[0]._tr.xpath("./w:trPr/w:tblHeader") if item.rows else []
                    )
                    has_header = any(
                        flag.get(
                            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val",
                            "true",
                        )
                        not in {"0", "false", "off"}
                        for flag in header_flags
                    )
                    blocks.extend(
                        _table_rows(
                            table_rows(item),
                            tuple(h[1] for h in headings),
                            f"DOCX 表 {table_number}",
                            has_header=has_header,
                        )
                    )
                    continue
                style = item.style.name if item.style else ""
                match = re.fullmatch(r"Heading ([1-9])", style)
                if match:
                    level = int(match[1])
                    while headings and headings[-1][0] >= level:
                        headings.pop()
                    headings.append((level, item.text))
                kind = f"h{min(int(match[1]), 6)}" if match else "p"
                if not match and (style.startswith("List") or item._p.xpath("./w:pPr/w:numPr")):
                    kind = "li"
                blocks.append(DocumentBlock(item.text, kind, tuple(h[1] for h in headings)))
            return _document(blocks)
        except ParsingFailed:
            raise
        except Exception as exc:
            raise ParsingFailed("DOCX_PARSE_FAILED") from exc


class ParserRegistry:
    def __init__(self) -> None:
        self._parsers: dict[str, Parser] = {
            "text/plain": TextParser(),
            "text/markdown": MarkdownParser(),
            "text/html": HtmlParser(),
            "application/pdf": PdfParser(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document": (
                DocxParser()
            ),
        }

    def parse(self, media_type: str, content: bytes) -> ParsedDocument:
        try:
            parser = self._parsers[media_type]
        except KeyError as exc:
            raise ParsingFailed("UNSUPPORTED_MEDIA_TYPE") from exc
        return parser.parse(content)
