"""Regression risks: lost table rules, split conditions, PDF provenance and citation drift."""

import re
from hashlib import sha256
from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from atlasmind.api.manual_structure import structured_document
from atlasmind.ingestion.chunking import DeterministicChunker
from atlasmind.ingestion.parsers import (
    DocxParser,
    MarkdownParser,
    ParsedDocument,
    ParsingFailed,
    PdfParser,
    TextParser,
)


def chunks(parsed, **kwargs):
    return DeterministicChunker(**kwargs).chunk(
        parsed, document_version_id="version", scope_ids=frozenset({"private"}), generation=2
    )


def pdf_bytes(page_lines):
    """Real PDF text objects with controllable positions, no external fixture generator."""
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    for lines in page_lines:
        page = writer.add_blank_page(width=600, height=800)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        stream = DecodedStreamObject()
        stream.set_data(
            "\n".join(
                f"BT /F1 12 Tf 1 0 0 1 {x} {y} Tm ({value}) Tj ET" for x, y, value in lines
            ).encode()
        )
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_markdown_inline_text_sections_and_qa_are_complete():
    raw = (
        "# 售后\n\n## 退货\n\n问\uff1a能否退货\uff1f\n\n"
        "答\uff1a支持 **七天** 退货\uff0c定制商品除外。\n\n## 运费\n\n由责任方承担。"
    ).encode()
    parsed = MarkdownParser().parse(raw)
    assert "问\uff1a能否退货\uff1f 答\uff1a支持 七天 退货\uff0c定制商品除外。" in parsed.paragraphs
    result = chunks(parsed)
    assert len(result) == 2
    assert "售后 > 退货" in result[0].content
    assert "定制商品除外" in result[0].content
    assert "售后 > 运费" in result[1].content


def test_markdown_table_rows_repeat_headers_and_keep_lists_in_order():
    parsed = MarkdownParser().parse(
        b"# Returns\n\n- Apply\n  - Attach proof\n- Wait\n\n"
        b"| Type | Days |\n| --- | --- |\n| Regular | 7 |\n| Custom | 0 |\n"
    )
    assert [b.text for b in parsed.blocks if b.kind == "li"] == ["Apply", "Attach proof", "Wait"]
    text = "\n".join(c.content for c in chunks(parsed))
    assert "Type\uff1aRegular | Days\uff1a7" in text
    assert "Type\uff1aCustom | Days\uff1a0" in text
    assert text.index("Attach proof") < text.index("Type\uff1aRegular")


def test_utf8_bom_wrapped_text_and_blank_line_qa():
    parsed = TextParser().parse(
        (
            "\ufeff# 规则\n\n问\uff1a如何申请\uff1f\n\n答\uff1a提交\n申请后等待审核。\n\n"
            "- 特殊情况请联系总部。"
        ).encode()
    )
    assert "问\uff1a如何申请\uff1f 答\uff1a提交申请后等待审核。" in parsed.paragraphs
    assert "规则" in chunks(parsed)[0].content
    english = TextParser().parse(b"A wrapped\nsentence.\n\nAnother paragraph.")
    assert english.paragraphs == ("A wrapped sentence.", "Another paragraph.")
    with pytest.raises(ParsingFailed, match="TEXT_ENCODING_UNSUPPORTED"):
        TextParser().parse(b"\xff")


def test_markdown_embedded_html_and_ordered_steps_do_not_lose_text():
    parsed = MarkdownParser().parse(
        b"# Rules\n\n3. Apply\n4. Wait\n\n<div>Policy <span>note</span></div>\n"
    )
    assert "3. Apply" in parsed.paragraphs and "4. Wait" in parsed.paragraphs
    assert "Policy note" in parsed.paragraphs


def test_docx_table_only_rules_nested_cells_and_document_order():
    doc = Document()
    doc.add_heading("Policy", 1)
    doc.add_paragraph("Before")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Type", "Days"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    table.cell(1, 0).text, table.cell(1, 1).text = "Custom", "0"
    nested = table.cell(1, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "No returns"
    doc.add_paragraph("After", style="List Bullet")
    output = BytesIO()
    doc.save(output)
    parsed = DocxParser().parse(output.getvalue())
    text = "\n".join(c.content for c in chunks(parsed))
    assert text.index("Before") < text.index("Type\uff1aCustom") < text.index("After")
    assert "Days\uff1a0 No returns" in text and "DOCX 表 1\uff0c行 2" in text
    assert parsed.blocks[-1].kind == "li"
    assert all("Policy" in c.content for c in chunks(parsed))


def test_table_only_docx_is_not_empty():
    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Refund within 7 days"
    output = BytesIO()
    doc.save(output)
    assert "Refund within 7 days" in DocxParser().parse(output.getvalue()).paragraphs


def test_unmarked_key_value_docx_table_does_not_invent_headers():
    doc = Document()
    doc.add_heading("Policy\nDetails", 1)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Fee", "0"
    table.cell(1, 0).text, table.cell(1, 1).text = "Days", "7"
    output = BytesIO()
    doc.save(output)
    result = chunks(DocxParser().parse(output.getvalue()))
    text = "\n".join(c.content for c in result)
    assert "列 1\uff1aDays | 列 2\uff1a7" in text
    assert "Fee\uff1aDays" not in text
    assert "Policy Details" in text
    for chunk in result:
        start, _, end = chunk.locator.removeprefix("paragraph=").partition("-")
        assert int(end or start) - int(start) + 1 == len(chunk.content.splitlines())


def test_pdf_removes_only_repeated_margin_text_and_preserves_page_provenance():
    data = pdf_bytes(
        [
            [
                (40, 780, "Company handbook"),
                (40, 650, "Common body rule"),
                (40, 620, f"Policy number {i}"),
                (40, 500, "Company handbook"),
                (40, 20, str(i)),
            ]
            for i in range(1, 4)
        ]
    )
    result = chunks(PdfParser().parse(data))
    text = "\n".join(c.content for c in result)
    assert text.count("Company handbook") == 3  # Identical body text must survive.
    assert text.count("Common body rule") == 3
    for number in range(1, 4):
        assert f"PDF 第 {number} 页" in text
        assert f"Policy number {number}" in text


def test_pdf_layout_preserves_column_boundaries():
    data = pdf_bytes(
        [[(40, 650, "Product"), (300, 650, "Days"), (40, 620, "Regular"), (300, 620, "7")]]
    )
    parsed = PdfParser().parse(data)
    assert any("Product | Days" in text for text in parsed.paragraphs)
    assert any("列 1\uff1aRegular | 列 2\uff1a7" in text for text in parsed.paragraphs)
    assert any("Product | Days" in c.content and "Regular" in c.content for c in chunks(parsed))


def test_pdf_blank_page_does_not_change_source_page_number():
    parsed = PdfParser().parse(pdf_bytes([[], [(40, 600, "Actual rule")]]))
    assert parsed.blocks[0].source.startswith("PDF 第 2 页")
    with pytest.raises(ParsingFailed, match="NO_EXTRACTABLE_TEXT"):
        PdfParser().parse(pdf_bytes([[]]))
    with pytest.raises(ParsingFailed, match="PDF_PARSE_FAILED"):
        PdfParser().parse(b"broken")


def test_pdf_scanned_page_is_not_silently_omitted(monkeypatch):
    class Page:
        class mediabox:
            bottom = 0
            top = 800

        images = (object(),)

        def extract_text(self, **kwargs):
            return ""

    class Reader:
        is_encrypted = False
        pages = (Page(),)

    monkeypatch.setattr("atlasmind.ingestion.parsers.PdfReader", lambda _: Reader())
    with pytest.raises(ParsingFailed, match="PDF_OCR_REQUIRED"):
        PdfParser().parse(b"scan")


def test_long_prose_is_bounded_without_losing_text_and_preserves_citation_contract():
    original = "普通商品支持七天退货\uff0c定制商品除外。" * 200
    result = chunks(TextParser().parse(original.encode()))
    assert len(result) > 1
    assert "".join(c.content for c in result) == original
    next_start = 1
    for chunk in result:
        assert len(chunk.content) <= 1500 and len(chunk.content.encode()) <= 4500
        assert chunk.content.endswith("。")
        assert chunk.content_digest == "sha256:" + sha256(chunk.content.encode()).hexdigest()
        assert chunk.scope_ids == frozenset({"private"}) and chunk.generation == 2
        match = re.fullmatch(r"paragraph=(\d+)(?:-(\d+))?", chunk.locator)
        assert int(match[1]) == next_start
        next_start += len(chunk.content.splitlines())
        assert int(match[2] or match[1]) == next_start - 1


def test_repeated_heading_context_and_unicode_byte_budget():
    parsed = MarkdownParser().parse(("# Policy\n\n## Returns\n\n" + "😀" * 800).encode())
    result = chunks(parsed, max_chars=1500, max_bytes=800)
    assert len(result) > 1
    assert all("Policy > Returns" in c.content and len(c.content.encode()) <= 800 for c in result)
    assert sum(c.content.count("😀") for c in result) == 800
    # Legacy callers can still construct ParsedDocument with its original two fields.
    assert chunks(ParsedDocument(("legacy",), "en"))[0].locator == "paragraph=1"


def test_new_markdown_rendering_matches_persisted_canonical_paragraphs():
    raw = b"# Policy\n\n- Keep **conditions** together\n\n<script>do_bad()</script>\n\n[Help](https://example.test)"
    result = chunks(MarkdownParser().parse(raw))
    paragraphs = [line for c in result for line in c.content.splitlines()]
    rendered = structured_document(raw, "text/markdown", paragraphs)

    def leaves(nodes):
        return [leaf for n in nodes for leaf in ([n] if "text" in n else leaves(n["children"]))]

    actual = leaves(rendered["nodes"])
    assert rendered["format"] == "markdown"
    assert [v["text"] for v in actual] == paragraphs
    assert [v["paragraph"] for v in actual] == list(range(1, len(paragraphs) + 1))
    assert "do_bad" not in str(rendered) and "https://" not in str(rendered)


def test_empty_and_oversize_heading_text_is_not_lost():
    parsed = MarkdownParser().parse(b"# Main\n\n## Empty\n\n## Full\n\nBody\n\n## Last")
    text = "\n".join(c.content for c in chunks(parsed))
    assert "Empty" in text and "Last" in text
    long_title = "Z" * 4000
    result = chunks(MarkdownParser().parse(("# " + long_title).encode()))
    assert all(len(c.content) <= 1500 for c in result)
    assert sum(c.content.split("\n", 1)[1].count("Z") for c in result) == 4000
